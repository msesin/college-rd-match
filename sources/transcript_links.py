"""Find and download the transcript links in a spreadsheet column. Plain Python, no AI.

A cell looks like (Airtable export):
    2023-2024.png (https://v5.airtableusercontent.com/...),2024-2025.JPG (https://...)
i.e. several `filename (url)` entries in one cell. Folder names come from the ID column,
never from the name column, which this module does not read.
"""

import csv
import hashlib
import mimetypes
import re
import time
from pathlib import Path
from urllib.parse import unquote, urlparse

import openpyxl
import requests

URL_PART = re.compile(r"\((https?://[^\s)]+)\)")
SAFE = re.compile(r"[^A-Za-z0-9._-]")
EXT_BY_TYPE = {"image/jpeg": ".jpg", "image/png": ".png", "application/pdf": ".pdf", "image/heic": ".heic",
               "image/webp": ".webp", "image/gif": ".gif"}


def parse_cell(text):
    """[(filename, url)] from one cell. The name is the text before each `(url)`, so names may contain brackets."""
    out, pos = [], 0
    text = text or ""
    for m in URL_PART.finditer(text):
        name = text[pos:m.start()].strip(" ,\n\t")
        if not name:
            name = unquote(Path(urlparse(m.group(1)).path).name)
        out.append((name, m.group(1)))
        pos = m.end()
    return out


def clean_folder(value):
    return SAFE.sub("_", str(value).strip())[:60] or "unknown"


def _column_index(letter):
    return openpyxl.utils.column_index_from_string(letter) - 1


def _iter_rows(path, sheet=None):
    """Rows (lists of cell values) after the header, from a .csv or .xlsx file."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        import sys

        csv.field_size_limit(sys.maxsize)
        with open(path, newline="", encoding="utf-8-sig") as fh:
            reader = csv.reader(fh)
            next(reader, None)
            yield from reader
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        yield from ws.iter_rows(min_row=2, values_only=True)


def read_rows(path, id_column, link_column, sheet=None):
    """[(row_number, id_value, [(filename, url)])]. Only these two columns are read."""
    id_i, link_i = _column_index(id_column), _column_index(link_column)
    rows = []
    for n, row in enumerate(_iter_rows(path, sheet), start=2):
        id_value = row[id_i] if len(row) > id_i else None
        cell = row[link_i] if len(row) > link_i else None
        rows.append((n, id_value if id_value not in (None, "") else f"row{n:04d}", parse_cell(str(cell or ""))))
    return rows


def unique_name(folder, name):
    p = folder / name
    stem, suffix, i = p.stem, p.suffix, 1
    while p.exists():
        i += 1
        p = folder / f"{stem}-{i}{suffix}"
    return p


def download(url, dest_folder, filename, session, timeout=60):
    """Returns (status, path or None, size, sha256). status: ok | expired | error:<detail>."""
    try:
        resp = session.get(url, timeout=timeout)
    except requests.RequestException as exc:
        return f"error:{type(exc).__name__}", None, 0, ""
    if resp.status_code in (401, 403, 410):
        return "expired", None, 0, ""
    if resp.status_code != 200:
        return f"error:http{resp.status_code}", None, 0, ""
    name = SAFE.sub("_", filename)
    if not Path(name).suffix:
        name += EXT_BY_TYPE.get(resp.headers.get("content-type", "").split(";")[0], mimetypes.guess_extension(
            resp.headers.get("content-type", "").split(";")[0]) or "")
    dest_folder.mkdir(parents=True, exist_ok=True)
    path = unique_name(dest_folder, name)
    path.write_bytes(resp.content)
    return "ok", path, len(resp.content), hashlib.sha256(resp.content).hexdigest()


def run(path, out_dir, id_column="A", link_column="C", sheet=None, limit=None, dry_run=False,
        delay=0.3, stop_after_expired=5, log=print):
    """Download every link. Returns a summary dict. Writes download_log.csv in out_dir."""
    rows = read_rows(path, id_column, link_column, sheet)
    if limit:
        rows = rows[:limit]
    out_dir = Path(out_dir)
    total_links = sum(len(r[2]) for r in rows)
    log(f"{len(rows)} rows, {sum(1 for r in rows if r[2])} with links, {total_links} links")
    if dry_run:
        return {"rows": len(rows), "links": total_links, "downloaded": 0, "expired": 0, "errors": 0}

    session = requests.Session()
    done = expired = errors = skipped = consecutive_expired = 0
    out_dir.mkdir(parents=True, exist_ok=True)
    log_path = out_dir / "download_log.csv"
    new_log = not log_path.exists()
    # Resume: files recorded as downloaded in an earlier run are skipped.
    already = set()
    if log_path.exists():
        with open(log_path, newline="") as fh:
            already = {(r["id"], r["filename"]) for r in csv.DictReader(fh) if r["status"] == "ok"}
    with open(log_path, "a", newline="") as fh:
        writer = csv.writer(fh)
        if new_log:
            writer.writerow(["id", "row", "filename", "status", "saved_as", "size_kb"])
        for n, id_value, links in rows:
            folder = out_dir / clean_folder(id_value)
            for filename, url in links:
                if (str(id_value), filename) in already:
                    skipped += 1
                    continue
                status, path, size, _ = download(url, folder, filename, session)
                writer.writerow([id_value, n, filename, status, path.name if path else "", round(size / 1024)])
                fh.flush()
                if status == "ok":
                    done += 1
                    consecutive_expired = 0
                elif status == "expired":
                    expired += 1
                    consecutive_expired += 1
                else:
                    errors += 1
                if consecutive_expired >= stop_after_expired:
                    log(f"Stopped: {consecutive_expired} links in a row were expired (or refused). "
                        "The links in this export have probably expired: export the table again and re-run.")
                    return {"rows": len(rows), "links": total_links, "downloaded": done, "expired": expired,
                            "errors": errors, "skipped": skipped, "stopped": True}
                time.sleep(delay)
    return {"rows": len(rows), "links": total_links, "downloaded": done, "expired": expired,
            "errors": errors, "skipped": skipped, "stopped": False}
